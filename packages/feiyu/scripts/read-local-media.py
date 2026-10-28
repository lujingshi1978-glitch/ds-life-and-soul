"""Bounded document and timeline reader. No macros, links or document code executes."""
import base64
import io
import json
import sys
from pathlib import Path


def image_record(image, label, detail="coarse", **extra):
    image = image.convert("RGB")
    image.thumbnail((1600, 1600) if detail == "fine" else (768, 768))
    stream = io.BytesIO()
    image.save(stream, "JPEG", quality=90 if detail == "fine" else 65)
    return {"label": label, "width": image.width, "height": image.height, "detail": detail, "mediaType": "image/jpeg", "data": base64.b64encode(stream.getvalue()).decode(), **extra}


def image_views(image, label, options, **extra):
    from PIL import ImageOps
    image = ImageOps.exif_transpose(image).convert("RGB")
    detail = options.get("detail", "auto")
    fine = detail == "fine" or (detail == "auto" and (options.get("textHint") or image.height > image.width * 2 or image.width > image.height * 3))
    mode = "fine" if fine else "coarse"
    views = [image_record(image, label + "·全图", mode, **extra)]
    if fine and max(image.size) > 1200:
        # Original-pixel crops, four per call. offset continues long screenshots.
        tile, stride = 1200, 1080
        boxes = [(x, y, min(x + tile, image.width), min(y + tile, image.height))
                 for y in range(0, image.height, stride) for x in range(0, image.width, stride)]
        start = max(0, int(options.get("tileOffset", 0)))
        for i, box in enumerate(boxes[start:start + 4], start):
            views.append(image_record(image.crop(box), f"{label}·局部{i + 1}", "fine", crop=list(box), **extra))
        views[0]["nextTileOffset"] = start + 4 if start + 4 < len(boxes) else None
        views[0]["totalTiles"] = len(boxes)
    return views


def read_file(file, options):
    suffix = file.suffix.lower()
    result = {"kind": suffix[1:], "text": "", "images": [], "partial": False}
    offset = max(0, int(options.get("offset", 0)))
    count = min(8, max(1, int(options.get("count", 4))))
    if suffix == ".pdf":
        import pymupdf
        with pymupdf.open(file) as document:
            if document.needs_pass:
                raise ValueError("PDF已加密，需要先解锁")
            result["pages"] = len(document)
            pages = range(min(offset, len(document)), min(offset + count, len(document)))
            for page_number in pages:
                page = document[page_number]
                text = page.get_text()
                result["text"] += f"\n[第{page_number + 1}页]\n{text[:12000]}"
                if len(text.strip()) < 80:
                    from PIL import Image
                    pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
                    result["images"].extend(image_views(Image.open(io.BytesIO(pix.tobytes("png"))), f"第{page_number + 1}页", {**options, "detail": "fine"}))
            result["nextOffset"] = offset + count if offset + count < len(document) else None
            result["partial"] = offset > 0 or result["nextOffset"] is not None
    elif suffix == ".docx":
        from docx import Document
        import zipfile
        with zipfile.ZipFile(file) as zipped:
            if sum(i.file_size for i in zipped.infolist()) > 128 * 1024 * 1024:
                raise ValueError("Word解压内容超过限制")
        document = Document(file)
        blocks = [p.text for p in document.paragraphs if p.text.strip()]
        blocks += [" | ".join(c.text for c in row.cells) for table in document.tables for row in table.rows]
        text = "\n".join(blocks)
        result.update(text=text[offset:offset + 30000], totalChars=len(text), nextOffset=offset + 30000 if len(text) > offset + 30000 else None, partial=offset > 0 or len(text) > 30000)
    elif suffix in {".gif", ".webp", ".png", ".jpg", ".jpeg"}:
        from PIL import Image
        Image.MAX_IMAGE_PIXELS = 24_000_000
        with Image.open(file) as animation:
            total = getattr(animation, "n_frames", 1)
            if total > 2000:
                raise ValueError("动图帧数超过限制")
            sample_count = min(count, total)
            samples = sorted(set(round(i * (total - 1) / max(1, sample_count - 1)) for i in range(sample_count)))
            at = 0
            for index in range(total):
                animation.seek(index)
                duration = max(1, int(animation.info.get("duration", 100)))
                if index in samples:
                    result["images"].extend(image_views(animation.copy(), f"帧{index}@{at}ms", options, index=index, timeMs=at, durationMs=duration))
                at += duration
            result.update(frames=total, durationMs=at, partial=total > len(samples), coverage="sampled-timeline" if total > 1 else "single-image")
    elif suffix in {".mp4", ".webm", ".mov", ".mkv", ".avi", ".m4v"}:
        import av
        start = max(0., float(options.get("start", 0)))
        with av.open(str(file)) as container:
            stream = container.streams.video[0]
            duration = float(container.duration / av.time_base) if container.duration else float(stream.duration * stream.time_base) if stream.duration else 120.
            end = min(duration, start + 120, float(options.get("end", start + 120)))
            if start >= end:
                raise ValueError("视频时间范围无效")
            targets = [start + (end - start) * (i + .5) / count for i in range(count)]
            for target in targets:
                container.seek(int(target / stream.time_base), stream=stream, backward=True)
                for frame in container.decode(stream):
                    timestamp = float(frame.time or 0)
                    if timestamp >= target:
                        result["images"].append(image_record(frame.to_image(), f"视频采样@{timestamp:.2f}s", timeSeconds=timestamp))
                        break
            result.update(durationSeconds=duration, window={"start": start, "end": end}, nextStart=end if end < duration else None, partial=True, coverage="sampled-frames-not-full-playback", hasAudio=bool(container.streams.audio))
    elif suffix in {".txt", ".md", ".csv", ".srt", ".vtt"}:
        text = file.read_text(encoding="utf-8-sig", errors="replace")
        result.update(text=text[offset:offset + 30000], totalChars=len(text), partial=len(text) > 30000 or offset > 0)
    else:
        raise ValueError("暂不支持该格式；旧版.doc请另存为.docx")
    all_images = result["images"]
    result["images"] = []
    size = 0
    for image in all_images:
        size += len(image["data"])
        if len(result["images"]) >= 12 or size > 8 * 1024 * 1024:
            result["visualBatchPartial"] = True
            result["partial"] = True
            break
        result["images"].append(image)
    result["text"] = result["text"][:40000]
    return result


if __name__ == "__main__":
    try:
        print(json.dumps({"ok": True, **read_file(Path(sys.argv[1]), json.loads(sys.argv[2]) if len(sys.argv) > 2 else {})}, ensure_ascii=False))
    except Exception as error:
        print(json.dumps({"ok": False, "error": str(error)[:200]}, ensure_ascii=False))
        sys.exit(1)
